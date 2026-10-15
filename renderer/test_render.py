"""Tests for the ok.json → xlsx renderer (issue #2).

Per ADR 0001 there are no layout tests, no vision, no pixel assertions:
every test here asserts the JSON→cell mapping (which contract field landed
in which графа) plus open/round-trip sanity of the produced file. Cell
coordinates always come from the FormLayout the renderer itself used —
never as literal geometry expectations.
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "contract"))

from ok_schema import validate_ok  # noqa: E402
from render import content_lines, render_ok  # noqa: E402
from template import form1_content_capacity  # noqa: E402

EXAMPLE = REPO / "examples" / "minimal" / "ok.json"
RENDER = Path(__file__).resolve().parent / "render.py"


def _load() -> dict:
    return json.loads(EXAMPLE.read_text(encoding="utf-8"))


def _two_operation_contract() -> dict:
    data = _load()
    op2 = json.loads(json.dumps(data["documents"][0]["operations"][0]))
    op2["number"] = 2
    op2["welds"] = ["Ш3"]
    op2["sketch_refs"] = ["КЭ2"]
    data["documents"][0]["operations"].append(op2)
    assert validate_ok(data) == []
    return data


# --- acceptance: every contract field lands in a defined графа ---------------


def test_every_contract_field_lands_in_a_defined_grapha() -> None:
    data = _load()
    wb, layout = render_ok(data)
    tags = {tag for tag, _ in content_lines(data)}
    assert {
        "product.process",
        "operations[0].number",
        "operations[0].position",
        "operations[0].transitions",
        "operations[0].transitions[0]",
        "operations[0].welds",
        "operations[0].sketch_refs",
        "operations[0].equipment",
        "operations[0].params",
        "documents[0].assumptions",
        "documents[0].assumptions[0]",
    } <= tags

    sheet = layout.sheets[0]
    ws = wb[sheet.title]
    assert ws[sheet.product_name_cell].value == data["product"]["name"]  # шапка
    assert ws[sheet.on_name_cell].value == data["product"]["name"]  # основная надпись
    assert ws[sheet.on_material_cell].value == data["product"]["material"]
    assert ws[sheet.table_number_cells[0]].value == data["documents"][0]["operations"][0]["number"]


def test_renderer_leaves_graphy_without_contract_fields_blank() -> None:
    data = _load()
    wb, layout = render_ok(data)
    ws = wb[layout.sheets[0].title]
    sheet = layout.sheets[0]
    # обозначение изделия/документа: no contract field → renderer invents nothing
    assert ws[sheet.product_designation_cell].value is None
    assert ws[sheet.on_designation_cell].value is None


# --- record lines per ГОСТ 3.1404-86 ----------------------------------------


def test_record_lines_for_minimal_example() -> None:
    data = _load()
    by_tag = dict(content_lines(data))
    assert by_tag["product.process"] == "Способ сварки: РД."
    assert by_tag["operations[0].number"] == "Операция 1"
    assert by_tag["operations[0].position"] == "Положение сварки: нижнее."
    assert by_tag["operations[0].transitions"] == "Переходы:"
    assert by_tag["operations[0].transitions[0]"] == "1. Сварить швы Ш1, Ш2 в нижнем положении. Шлак и брызги удалить."
    assert by_tag["operations[0].welds"] == "Швы: Ш1, Ш2."
    assert by_tag["operations[0].sketch_refs"] == "Ссылки на эскизы: КЭ1."
    assert by_tag["operations[0].equipment"] == "Электроды: УОНИ-13/55."
    assert by_tag["operations[0].params"] == (
        "Режим: постоянный ток, обратная полярность, 130 А; электрод Ø4 мм; 2 прохода."
    )
    assert by_tag["documents[0].assumptions"] == "Примечания:"
    assert by_tag["documents[0].assumptions[0]"] == data["documents"][0]["assumptions"][0]


def test_ac_current_record_omits_polarity() -> None:
    data = _load()
    params = data["documents"][0]["operations"][0]["params"]
    params["current_type"] = "переменный"
    params.pop("polarity")
    # contract/README.md is normative: polarity is omitted for "переменный".
    # NOTE: validate_ok() currently rejects this shape (ok_schema._check_keys
    # requires the polarity key unconditionally) — a known contract-validator
    # discrepancy, issue #1 scope, not touched here. The renderer follows the
    # README: it includes polarity only when the contract carries it.
    mode = dict(content_lines(data))["operations[0].params"]
    assert mode.startswith("Режим: переменный ток,")
    assert "полярность" not in mode


def test_content_lines_fill_form1_slots_in_order() -> None:
    data = _load()
    wb, layout = render_ok(data)
    lines = content_lines(data)
    ws = wb[layout.sheets[0].title]
    for (_, text), (title, cell) in zip(lines, layout.content_slots()):
        assert title == layout.sheets[0].title
        assert ws[cell].value == text


# --- pagination onto форма 1а ------------------------------------------------


def test_overflow_paginates_onto_form_1a() -> None:
    data = _two_operation_contract()
    lines = content_lines(data)
    n_ops = len(data["documents"][0]["operations"])
    cap1 = form1_content_capacity(n_ops)
    assert 1 < len(lines) - cap1, "scenario must actually overflow форма 1"

    wb, layout = render_ok(data)
    assert [sheet.title for sheet in layout.sheets] == wb.sheetnames
    assert len(layout.sheets) == 2
    assert layout.sheets[1].title.startswith("Форма 1а")

    ws1, ws2 = wb[layout.sheets[0].title], wb[layout.sheets[1].title]
    for i, (_, text) in enumerate(lines[:cap1]):
        assert ws1[layout.sheets[0].content_cells[i]].value == text
    for i, (_, text) in enumerate(lines[cap1:]):
        assert ws2[layout.sheets[1].content_cells[i]].value == text

    sheet2 = layout.sheets[1]
    assert ws2[sheet2.sheet_no_cell].value == 2  # «Лист»
    assert ws2[sheet2.sheets_total_cell].value == 2  # «Листов»


# --- open / round-trip sanity of the produced file ----------------------------


def test_rendered_file_round_trips(tmp_path) -> None:
    data = _load()
    wb, layout = render_ok(data)
    out = tmp_path / "ok.xlsx"
    wb.save(out)

    again = openpyxl.load_workbook(out)
    assert again.sheetnames == [sheet.title for sheet in layout.sheets]
    ws = again[layout.sheets[0].title]
    sheet = layout.sheets[0]
    assert ws[sheet.product_name_cell].value == data["product"]["name"]
    assert ws[sheet.on_material_cell].value == data["product"]["material"]
    assert ws[sheet.content_cells[0]].value == content_lines(data)[0][1]


# --- CLI ----------------------------------------------------------------------


def test_cli_renders_valid_contract(tmp_path) -> None:
    out = tmp_path / "ok.xlsx"
    proc = subprocess.run(
        [sys.executable, str(RENDER), str(EXAMPLE), str(out)],
        capture_output=True, text=True,
    )
    assert proc.returncode == 0, proc.stderr
    assert out.exists()
    assert openpyxl.load_workbook(out).sheetnames == ["Форма 1"]


def test_cli_refuses_invalid_contract(tmp_path) -> None:
    data = _load()
    del data["product"]["material"]
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    proc = subprocess.run(
        [sys.executable, str(RENDER), str(bad)],
        capture_output=True, text=True,
    )
    assert proc.returncode == 1
    assert "product.material" in proc.stderr


def test_cli_usage_error() -> None:
    proc = subprocess.run([sys.executable, str(RENDER)], capture_output=True, text=True)
    assert proc.returncode == 2
