"""ok.json → ok.xlsx renderer (ADR 0001, issue #2).

Deterministic mapping of a *valid* ok.json contract into the ГОСТ 3.1118-82
форма 1/1а workbook built by template.py. The renderer makes zero engineering
decisions: it never invents, infers or rewrites content — every value it
places comes verbatim from the contract, composed into ГОСТ 3.1404-86 record
lines by pure formatting (joining, Russian plural forms, number formatting).
Graphы the contract has no field for (обозначение документа, коды, times,
signatures) stay blank and are filled by hand at приёмка.

Field → графа mapping (normative for tests; see renderer/README.md):

- product.name      → шапка «Наименование изделия» + основная надпись (наименование)
- product.material  → основная надпись, графа «материал»
- product.process   → содержание, «Способ сварки: …»
- operations[i].number → таблица, графа «Номер операции» + содержание «Операция N»
- operations[i].position / transitions / welds / sketch_refs / equipment / params
                    → содержание операции (one record line each)
- documents[0].assumptions[k] → содержание, block «Примечания:»

CLI (validates through the contract validator first — the pipeline refuses to
render an invalid contract; rendering itself still checks nothing):

    python3 renderer/render_ok.py <ok.json> [out.xlsx]
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook

from template import FormLayout, build_document

_CONTRACT_DIR = Path(__file__).resolve().parents[1] / "contract"
sys.path.insert(0, str(_CONTRACT_DIR))

from ok_schema import validate_ok  # noqa: E402  (path set above)


def _fmt_num(value: float) -> str:
    """130 → '130', 4.5 → '4.5' (round-trips the JSON number, no units here)."""
    return f"{value:g}"


def _plural_ru(n: int, one: str, few: str, many: str) -> str:
    """Russian plural forms — pure formatting (2 прохода, 5 проходов)."""
    n100 = abs(n) % 100
    if 11 <= n100 <= 14:
        return many
    n10 = n100 % 10
    if n10 == 1:
        return one
    if 2 <= n10 <= 4:
        return few
    return many


def _mode_text(params: dict) -> str:
    """Режим сварки as one ГОСТ 3.1404-86-style record line.

    Polarity is present only for постоянный ток (the contract omits it
    otherwise); the renderer merely includes what is there.
    """
    parts = [f"{params['current_type']} ток"]
    if "polarity" in params:
        parts.append(f"{params['polarity']} полярность")
    parts.append(f"{_fmt_num(params['current_a'])} А")
    passes = params["pass_count"]
    plural = _plural_ru(passes, "проход", "прохода", "проходов")
    return (
        f"Режим: {', '.join(parts)}; "
        f"электрод Ø{_fmt_num(params['electrode_diameter_mm'])} мм; {passes} {plural}."
    )


def content_lines(data: dict) -> list[tuple[str, str]]:
    """Contract → record lines for the зона содержания операции.

    Returns (tag, text) pairs in fill order. ``tag`` is the contract path the
    line renders (validator-style: ``operations[0].params``), so tests can
    prove every contract field lands in a defined графа. Line composition is
    formatting only: fixed labels + contract values, nothing inferred.
    """
    lines: list[tuple[str, str]] = [
        ("product.process", f"Способ сварки: {data['product']['process']}.")
    ]
    document = data["documents"][0]
    for i, operation in enumerate(document["operations"]):
        p = f"operations[{i}]"
        lines.append((f"{p}.number", f"Операция {operation['number']}"))
        lines.append((f"{p}.position", f"Положение сварки: {operation['position']}."))
        lines.append((f"{p}.transitions", "Переходы:"))
        for j, transition in enumerate(operation["transitions"]):
            lines.append((f"{p}.transitions[{j}]", f"{transition['number']}. {transition['description']}"))
        lines.append((f"{p}.welds", "Швы: " + ", ".join(operation["welds"]) + "."))
        lines.append((f"{p}.sketch_refs", "Ссылки на эскизы: " + ", ".join(operation["sketch_refs"]) + "."))
        lines.append((f"{p}.equipment", f"Электроды: {operation['equipment']['electrode_type']}."))
        lines.append((f"{p}.params", _mode_text(operation["params"])))
    lines.append(("documents[0].assumptions", "Примечания:"))
    for k, assumption in enumerate(document["assumptions"]):
        lines.append((f"documents[0].assumptions[{k}]", assumption))
    return lines


def render_ok(data: dict) -> tuple[Workbook, FormLayout]:
    """Render a valid contract into a workbook plus its layout descriptor."""
    lines = content_lines(data)
    operations = data["documents"][0]["operations"]
    wb = Workbook()
    wb.remove(wb.active)
    layout = build_document(wb, len(operations), len(lines))

    sheet1 = layout.sheets[0]
    ws = wb[sheet1.title]
    ws[sheet1.product_name_cell] = data["product"]["name"]
    ws[sheet1.on_name_cell] = data["product"]["name"]
    ws[sheet1.on_material_cell] = data["product"]["material"]
    for operation, cell in zip(operations, sheet1.table_number_cells):
        ws[cell] = operation["number"]

    slots = layout.content_slots()
    if len(slots) < len(lines):
        raise RuntimeError("renderer built fewer content slots than the contract needs")
    for (_, text), (sheet_title, cell) in zip(lines, slots):
        wb[sheet_title][cell] = text
    return wb, layout


def main(argv: list[str]) -> int:
    args = argv[1:]
    if not args or args[0] in ("-h", "--help") or len(args) > 2:
        print(__doc__.strip())
        return 0 if args and args[0] in ("-h", "--help") else 2

    path = Path(args[0])
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except OSError as exc:
        print(f"error: cannot read {path}: {exc}", file=sys.stderr)
        return 2
    except json.JSONDecodeError as exc:
        print(f"error: {path} is not valid JSON: line {exc.lineno}, column {exc.colno}: {exc.msg}", file=sys.stderr)
        return 1

    errors = validate_ok(data)
    if errors:
        for error in errors:
            print(f"error: {error}", file=sys.stderr)
        print(f"invalid: {len(errors)} error(s) in {path}; refusing to render", file=sys.stderr)
        return 1

    out = Path(args[1]) if len(args) == 2 else path.with_suffix(".xlsx")
    wb, layout = render_ok(data)
    wb.save(out)
    print(
        f"rendered {out}: {data['product']['name']!r}, "
        f"{len(layout.sheets)} sheet(s) (форма 1{' + форма 1а' if len(layout.sheets) > 1 else ''}), "
        f"{len(content_lines(data))} content line(s)"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
